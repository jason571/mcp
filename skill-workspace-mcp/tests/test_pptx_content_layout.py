"""Content helper regressions: editable CJK and dense layouts, not only ZIP validity."""
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest

from pptx import Presentation

scripts = Path(__file__).resolve().parents[3] / 'skills/pptx/scripts'
sys.path.insert(0, str(scripts))
spec = importlib.util.spec_from_file_location('pptx_content_renderer', scripts / 'render_deck.py')
renderer = importlib.util.module_from_spec(spec)
spec.loader.exec_module(renderer)


def fixture():
    return {'title': '小型场馆运营改进', 'slides': [
        {'layout': 'cover', 'title': '小型场馆运营改进：先稳定交付，再扩大经营',
         'subtitle': '定性分析｜聚焦获客、体验与成本，未使用未经核实的市场数字'},
        {'layout': 'cards', 'title': '经营改进不能只看上座率：应同时验证体验、复购与成本',
         'subtitle': '将三个观察转化为可检查的经营动作，而不是笼统增加营销投入',
         'items': [
             {'heading': '获客：从渠道追到复购', 'text': '记录客源渠道与首次到店动机，再对比复购表现。高曝光不代表有效获客，应按真实成交和回访反馈调整投放。'},
             {'heading': '体验：消除高频摩擦', 'text': '逐项观察预约、到店、使用设备和离场环节。先修复排队、操作不清和内容单调等可验证问题，再增加花哨的新功能。'},
             {'heading': '成本：看单位服务效率', 'text': '区分固定投入与随客流变化的支出。把人员排班、能耗和设备故障纳入同一复盘，避免节省费用却降低核心体验。'}],
         'note': '建议以同一观察周期记录变化；以上为定性方法，不是收益承诺。'},
        {'layout': 'comparison', 'title': '先改善核心体验，比直接扩张更容易核验效果',
         'columns': ['比较维度', '优先改善体验', '直接扩大规模'],
         'rows': [['投入重点', '处理已知服务摩擦，验证客户是否愿意再次到店', '增加空间、设备或渠道，需要更高固定投入'],
                  ['主要风险', '只修局部问题而没有追踪复购，难以证明真实价值', '在需求未经验证时扩张，闲置产能会放大现金流压力'],
                  ['验证方式', '结合回访、故障记录与复购趋势做阶段复盘', '核验稳定客源、人员供给与现金流缓冲']],
         'note': '不能仅凭短期客流增加决定扩张；比较真实业务约束，不假设固定增长比例。'},
        {'layout': 'process', 'title': '用三个阶段把改进变成可检查的行动',
         'items': [{'heading': '建立基线', 'text': '整理客流、反馈、设备故障和服务成本，确定一个最值得处理的问题。'},
                   {'heading': '小范围试行', 'text': '明确负责人和观察周期，先在可控场次验证服务改进，记录实际变化。'},
                   {'heading': '复盘再推广', 'text': '检查体验是否改善及成本是否可承受；证据不足时修正方案，不急于扩张。'}],
         'note': '下一步：选择一个高频摩擦点，设定负责人、观察周期和可核验的成功条件。'}]}


class ContentLayoutTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.output = Path(self.tmp.name) / 'editable.pptx'

    def test_dense_four_slide_fixture_is_editable_and_retains_text(self):
        data = fixture()
        result = renderer.render(data, self.output)
        deck = Presentation(self.output)
        self.assertEqual(result['slides'], 4)
        self.assertEqual(json.loads(self.output.with_name('.editable.pptx.content.json').read_text()), data)
        self.assertEqual(len(deck.slides), 4)
        all_text = ''.join(s.text.replace('\n', '') for slide in deck.slides for s in slide.shapes if s.has_text_frame)
        for slide in data['slides']:
            self.assertIn(slide['title'], all_text)
            for item in slide.get('items', []):
                self.assertIn(item['heading'], all_text)
                self.assertIn(item['text'], all_text)
        for slide in deck.slides:
            self.assertFalse(any(s.shape_type == 13 for s in slide.shapes))
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            self.assertEqual(run.font.name, 'Microsoft YaHei')
                            self.assertEqual(run._r.rPr.get('lang'), 'zh-CN')

    def test_oversized_copy_rejects_without_overwriting_previous_delivery(self):
        data = fixture()
        self.output.write_bytes(b'previous accepted file')
        data['slides'][1]['items'][0]['text'] = '不能删除的原始用户内容' * 300
        with self.assertRaisesRegex(renderer.ContentError, r'slides\[1\].items'):
            renderer.render(data, self.output)
        self.assertEqual(self.output.read_bytes(), b'previous accepted file')
        saved = json.loads(self.output.with_name('.editable.pptx.content.json').read_text())
        self.assertEqual(saved, data)

    def test_notes_and_wrapped_title_regions_do_not_overlap(self):
        composer = renderer.Composer(fixture())
        composer.compose()
        # Every page is checked internally. Inspect final slide's geometric gap.
        regions = {v[4]: v for v in composer.regions}
        note = regions['slides[3].note']
        for key, (_, y, _, h, _) in regions.items():
            if '.items[' in key:
                self.assertLess(y + h, note[1])

    def test_wrap_keeps_copy_and_avoids_orphan_cjk_punctuation(self):
        composer = renderer.Composer(fixture())
        for value in ('经营改进不能只看上座率：应同时验证体验、复购与成本',
                      '记录客源渠道与首次到店动机，再对比复购表现。高曝光不代表有效获客，应按真实成交和回访反馈调整投放。'):
            wrapped, _ = composer.wrap(value, 3.0, 18, 'test')
            self.assertEqual(wrapped.replace('\n', ''), value)
            self.assertFalse(any(line.startswith(('。', '，', '！')) for line in wrapped.split('\n')))

    def test_unknown_fields_and_encoded_arrays_are_not_guessed(self):
        data = fixture()
        data['slides'][1]['items'] = json.dumps(data['slides'][1]['items'])
        with self.assertRaisesRegex(renderer.ContentError, 'nonempty array'):
            renderer.render(data, self.output)
        data = fixture()
        data['slides'][0]['accidental'] = True
        with self.assertRaisesRegex(renderer.ContentError, 'unexpected slide fields'):
            renderer.render(data, self.output)

    def test_no_implicit_slides_or_assistant_branding(self):
        data = {'slides': [{'layout': 'rows', 'title': '自定义主题',
                           'items': [{'heading': '关键依据', 'text': '仅保留用户需要的说明。'}]}]}
        renderer.render(data, self.output)
        deck = Presentation(self.output)
        self.assertEqual(len(deck.slides), 1)
        self.assertEqual(deck.core_properties.author, '')
        self.assertNotIn('CLHH', ''.join(s.text for s in deck.slides[0].shapes if s.has_text_frame))

    def test_item_notes_are_preserved_only_in_speaker_notes(self):
        data = fixture()
        data['slides'][1]['notes'] = '原有讲解备注'
        data['slides'][1]['items'][0]['notes'] = '补充材料不占卡片区域'
        before = json.dumps(data, ensure_ascii=False)
        renderer.render(data, self.output)
        self.assertEqual(json.dumps(data, ensure_ascii=False), before)
        slide = Presentation(self.output).slides[1]
        self.assertIn('原有讲解备注', slide.notes_slide.notes_text_frame.text)
        self.assertIn('补充材料不占卡片区域', slide.notes_slide.notes_text_frame.text)
        self.assertNotIn('补充材料不占卡片区域', ''.join(s.text for s in slide.shapes if s.has_text_frame))
        data['slides'][1]['items'][0]['notes'] = ['not a string']
        with self.assertRaisesRegex(renderer.ContentError, r'items\[0\].notes'):
            renderer.render(data, self.output)
        data = fixture()
        data['slides'][1]['items'][0]['accidental'] = True
        with self.assertRaisesRegex(renderer.ContentError, r'items\[0\].*accidental'):
            renderer.render(data, self.output)

    def test_invalid_comparison_and_broken_unicode_are_local_errors(self):
        data = fixture()
        data['slides'][2]['rows'][0].pop()
        with self.assertRaisesRegex(renderer.ContentError, 'match the column count'):
            renderer.render(data, self.output)
        data = fixture()
        data['slides'][1]['items'][0]['text'] = '增长\ufffd'
        with self.assertRaisesRegex(renderer.ContentError, 'damaged text encoding'):
            renderer.render(data, self.output)

    def test_page_count_is_not_a_blocking_condition(self):
        data = fixture()
        data['slides'].append(data['slides'][-1])
        result = renderer.render(data, self.output, expected_slides=4)
        self.assertGreaterEqual(result['slides'], 5)
        self.assertTrue(self.output.exists())

    def test_dense_comparison_continues_complete_rows_without_losing_cells(self):
        data = fixture()
        original = data['slides'][2]
        original['rows'] = original['rows'] * 3
        before = json.dumps(data, ensure_ascii=False)
        result = renderer.render(data, self.output)
        self.assertGreater(result['slides'], 4)
        self.assertEqual(json.dumps(data, ensure_ascii=False), before)
        deck = Presentation(self.output)
        text = ''.join(s.text.replace('\n', '') for slide in deck.slides for s in slide.shapes if s.has_text_frame)
        self.assertEqual(text.count(original['rows'][0][1]), 3)

    def test_wide_comparison_retains_every_column_and_row_identity(self):
        for count in (5, 8):
            with self.subTest(columns=count):
                columns = ['比较项目'] + [f'指标{i}' for i in range(1, count)]
                rows = [[f'项目{r}'] + [f'项目{r}的值{i}' for i in range(1, count)]
                        for r in range(3)]
                data = {'slides': [{'title': '完整宽表', 'layout': 'comparison',
                                    'columns': columns, 'rows': rows}]}
                before = json.dumps(data, ensure_ascii=False)
                renderer.render(data, self.output)
                self.assertEqual(json.dumps(data, ensure_ascii=False), before)
                deck = Presentation(self.output)
                texts = [s.text.replace('\n', '') for slide in deck.slides
                         for s in slide.shapes if s.has_text_frame]
                for column in columns[1:]:
                    self.assertEqual(texts.count(column), 1)
                for row in rows:
                    for cell in row[1:]:
                        self.assertEqual(texts.count(cell), 1)
                    self.assertEqual(texts.count(row[0]), (count - 2) // 3 + 1)

    def test_wide_comparison_invalid_row_rejects_without_hiding_cells(self):
        data = {'slides': [{'title': '不合法表', 'layout': 'comparison',
                            'columns': ['标识', '一', '二', '三', '四'],
                            'rows': [['行', '值', '值', '值']]}]}
        with self.assertRaisesRegex(renderer.ContentError, 'match the column count'):
            renderer.render(data, self.output)


if __name__ == '__main__':
    unittest.main()
