import base64
import copy
import hashlib
import json
import os
import shutil
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
import zipfile
from pathlib import Path

from skill_workspace.common import package_path, relative, validate_request, write, read
from skill_workspace.broker import container_args, run_job, docker, finish, runtime_inventory, readable_package_snapshot
from skill_workspace.workspaces import Workspaces


def request(code="print('ok')", runtime="python", outputs=None):
    return {"package_name": "unrelated-guide", "package_digest": "a" * 64, "runtime": runtime,
            "files": [{"path": "main.py" if runtime == "python" else "main.js", "content": code}],
            "entrypoint": "main.py" if runtime == "python" else "main.js", "outputs": outputs or ["output.txt"]}


class SecurityTests(unittest.TestCase):
    def test_artifact_identity_tracks_workspace_path_revisions_not_filenames(self):
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"WORKSPACE_IMAGE_ID":"sha256:" + "a" * 64}):
            root = Path(directory)
            content = b"fixture"
            identities = []
            for index, (workspace, path, revision) in enumerate([
                ("ws_a", "chart.png", 1), ("ws_a", "chart.png", 2),
                ("ws_b", "chart.png", 1), ("ws_a", "other/chart.png", 1),
            ]):
                job = root / str(index)
                job.mkdir()
                finish(job, {"workspace_id":workspace,"workspace_revision":revision,"outputs":[path]},
                    json.dumps({"exit_code":0,"files":[{"path":path,"data":base64.b64encode(content).decode(),
                        "size_bytes":len(content),"sha256":hashlib.sha256(content).hexdigest()}]}))
                artifact = read(job / "result.json")["files"][0]
                identity = artifact["_meta"]["exoharness"]["artifact_identity"]
                self.assertEqual(identity["revision"], revision + 1)
                self.assertNotIn(workspace, identity["key"])
                identities.append(identity)
            self.assertEqual(identities[0]["key"], identities[1]["key"])
            self.assertNotEqual(identities[0]["key"], identities[2]["key"])
            self.assertNotEqual(identities[0]["key"], identities[3]["key"])
            self.assertEqual((root / "0/published/chart.png").read_bytes(), content, "superseded files remain available as evidence")

    def test_container_failure_never_falls_back_to_host_python_or_shell(self):
        for runtime in ["python", "shell"]:
            with self.subTest(runtime=runtime), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                image = "sha256:" + "a" * 64
                workspace = Workspaces(root).create(runtime_image=image)
                write(root / "worker.json", {"image": image})
                write(root / "task-code.json", {"enabled": True, "image": image})
                marker = root / "must-not-be-written-on-host"
                code = f"from pathlib import Path\nPath({str(marker)!r}).write_text('escaped')" if runtime == "python" else f"printf escaped > '{marker}'"
                job = root / "jobs" / ("work_" + "b" * 64)
                write(job / "request.json", {"workspace_id": workspace["workspace_id"], "workspace_revision": 0,
                    "runtime_image": image, "runtime": runtime, "package_name": None, "package_digest": None,
                    "entrypoint": "main.py" if runtime == "python" else "script.sh",
                    "files": [{"path": "main.py" if runtime == "python" else "script.sh", "content": code}], "outputs": []})
                with patch("skill_workspace.broker.inspect", return_value=None), patch("skill_workspace.broker.docker", side_effect=RuntimeError("sandbox unavailable")) as engine:
                    with self.assertRaisesRegex(RuntimeError, "sandbox unavailable"):
                        run_job(job, root / "packages", image, root / "worker.json")
                    self.assertEqual(engine.call_args.args[0], "create")
                    self.assertEqual(engine.call_count, 1)
                self.assertFalse(marker.exists())

    def test_traversal_duplicate_and_unapproved_profiles(self):
        for path in ["../secret", "/tmp/a", "a/../b", "a//b", "a\\b", ".", "a/./b"]:
            with self.assertRaises(ValueError):
                relative(path)
        for mutate in [lambda x: x.update(runtime="host_shell"), lambda x: x.update(entrypoint="/bin/bash"),
                       lambda x: x.update(package_name="../../host"), lambda x: x.update(outputs=["a.exe"]),
                       lambda x: x["files"].append(x["files"][0]), lambda x: x.update(outputs=["a/output.txt", "b/output.txt"])]:
            value = request()
            mutate(value)
            with self.assertRaises(ValueError):
                validate_request(value)
        validate_request(request())
        validate_request(request(outputs=["index.html"]))  # Downloadable; inline preview is separately sandboxed.

    def test_container_policy_cannot_be_relaxed_by_request(self):
        args = container_args(Path("/approved/job"), Path("/approved/skill"), "sha256:" + "a" * 64, Path("/approved/runner.py"))
        self.assertIn("none", args)
        self.assertIn("--read-only", args)
        self.assertIn("65532:65532", args)
        self.assertNotIn("--privileged", args)
        self.assertFalse(any("docker.sock" in arg for arg in args))
        mounts = [args[i + 1] for i, arg in enumerate(args) if arg == "--mount"]
        self.assertTrue(all(mount.endswith(",readonly") for mount in mounts))
        with self.assertRaises(ValueError):
            container_args(Path("/job"), Path("/skill"), "arbitrary:latest", Path("/runner"))
        browser = container_args(Path("/approved/job"), None, "sha256:"+"a"*64, Path("/approved/browser_runner.py"), browser=True)
        self.assertEqual(args[args.index("--pids-limit")+1], "128")
        self.assertEqual(browser[browser.index("--pids-limit")+1], "256")
        self.assertIn("/tmp:rw,noexec,nosuid,size=128m,mode=1777", args)
        self.assertIn("/tmp:rw,noexec,nosuid,size=512m,mode=1777", browser)
        self.assertEqual(browser[browser.index("--memory")+1], "2g")
        self.assertIn("LP_NUM_THREADS=2", browser)
        self.assertIn("--read-only", browser)
        self.assertEqual(browser[browser.index("--network")+1], "none")
        self.assertNotIn("--privileged", browser)

    def test_symlink_package_and_artifact_tampering(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            actual = root / "real"
            actual.mkdir()
            (actual / "SKILL.md").write_text("private")
            (root / "unrelated-guide").symlink_to(actual)
            with self.assertRaises(ValueError):
                package_path(root, "unrelated-guide", "a" * 64)
            value = {"exit_code": 0, "files": [{"path": "output.txt", "size_bytes": 5, "sha256": "0" * 64, "data": base64.b64encode(b"wrong").decode()}]}
            with self.assertRaisesRegex(ValueError, "integrity"):
                finish(root, request(), json.dumps(value).encode())

    def test_package_snapshot_rejects_nested_links(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            package = root / "package"
            package.mkdir()
            (root / "private").write_text("must not be mounted")
            (package / "secret.txt").symlink_to(root / "private")
            with self.assertRaisesRegex(ValueError, "links and special"):
                readable_package_snapshot(package, root / "snapshot")
            self.assertFalse((root / "snapshot").exists())


@unittest.skipUnless(os.getenv("SKILL_WORKSPACE_INTEGRATION") == "1", "requires local Docker and built immutable image")
class DockerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="skill-workspace-test-")
        self.root = Path(self.temp.name)
        self.job = self.root / "jobs" / ("work_" + os.urandom(32).hex())
        self.job.mkdir(parents=True)
        write(self.root / "packages.json", {"unrelated-guide": ["a" * 64]})
        self.package = self.root / "packages" / "unrelated-guide" / ("a" * 64)
        self.package.mkdir(parents=True)
        (self.package / "SKILL.md").write_text("original skill instructions")

    def tearDown(self):
        docker("rm", "-f", "exoharness-" + self.job.name, check=False)
        self.temp.cleanup()

    def prepare(self, data):
        validate_request(data)
        write(self.job / "request.json", data)
        write(self.job / "input/request.json", data)
        (self.job / "input/request.json").chmod(0o644)
        write(self.job / "state.json", {"status": "queued"})

    def run_job(self):
        run_job(self.job, self.root / "packages", os.environ["WORKSPACE_IMAGE_ID"], self.root / "worker.json")

    def test_real_python_artifacts_and_security_boundaries(self):
        code = """import os, socket
from pathlib import Path
from pptx import Presentation
from docx import Document
assert not Path('/var/run/docker.sock').exists()
assert not Path('/run/secrets').exists()
assert os.getuid() == 65532
try:
 Path('/skill/SKILL.md').write_text('bad')
 raise AssertionError('read-only mount failed')
except OSError: pass
try:
 socket.create_connection(('1.1.1.1', 443), timeout=1)
 raise AssertionError('network isolation failed')
except OSError: pass
p = Presentation()
s = p.slides.add_slide(p.slide_layouts[0])
s.shapes.title.text = '娱乐行业分析 - 执行验证'
p.save('report.pptx')
d = Document()
d.add_paragraph('同一个通用执行工具，也可以输出其他文档')
d.save('report.docx')
"""
        self.prepare(request(code, outputs=["report.pptx", "report.docx"]))
        self.run_job()
        self.assertEqual(read(self.job / "state.json")["status"], "completed")
        for filename, required in [("report.pptx", "ppt/presentation.xml"), ("report.docx", "word/document.xml")]:
            with zipfile.ZipFile(self.job / "published" / filename) as archive:
                self.assertIn(required, archive.namelist())
                self.assertIsNone(archive.testzip())

    def test_real_node_pptx_and_libreoffice_render(self):
        original = Path(__file__).resolve().parents[3] / "skills/pptx"
        if not original.is_dir():
            original = Path(__file__).resolve().parents[3] / "examples/skills/pptx"
        shutil.copytree(original, self.package, dirs_exist_ok=True)
        code = """const pptxgen = require('pptxgenjs');
const {execFileSync} = require('child_process');
(async () => { const p = new pptxgen(); const s = p.addSlide();
s.addText('Skill-owned presentation', {x:1,y:1,w:7,h:1,fontSize:28});
await p.writeFile({fileName:'report.pptx'});
for (const args of [
 ['/skill/scripts/office/validate.py','report.pptx'],
 ['/skill/scripts/office/soffice.py','--headless','--convert-to','pdf','--outdir','out','report.pptx'],
 ['/skill/scripts/thumbnail.py','report.pptx','preview']
]) console.log(execFileSync('python', args, {timeout:60000}).toString());
})();"""
        self.prepare(request(code, "node", ["report.pptx", "out/report.pdf", "preview.jpg"]))
        self.run_job()
        self.assertTrue((self.job / "published/report.pdf").read_bytes().startswith(b"%PDF"))
        self.assertTrue((self.job / "published/preview.jpg").read_bytes().startswith(b"\xff\xd8"))
        self.assertIn("All validations PASSED", read(self.job / "execution.json")["log"])

    def test_real_pptx_text_options_bug_is_caught_and_only_deck_is_published(self):
        original = Path(__file__).resolve().parents[3] / "skills/pptx"
        shutil.copytree(original, self.package, dirs_exist_ok=True)
        code = """const PptxGenJS = require('pptxgenjs');
const {spawnSync} = require('child_process');
const assert = require('assert');
async function build(file, helper) {
 const p = new PptxGenJS(); p.layout = 'LAYOUT_WIDE';
 p.theme = {headFontFace:'Noto Sans CJK SC',bodyFontFace:'Noto Sans CJK SC',lang:'zh-CN'};
 const s=p.addSlide();
 for(let i=0;i<4;i++) s.addText('中文分析结论 '+i,helper({x:0.7,y:0.8+i,w:10,h:0.7,fontSize:22,color:'172A3A'}));
 await p.writeFile({fileName:file});
}
function bad(s,options){return Object.assign({fontFace:'Noto Sans CJK SC'},options||{});}
function good(options){return {...options,fontFace:'Noto Sans CJK SC',margin:0};}
(async()=>{
 await build('broken.pptx',bad);
 const failed=spawnSync('python',['/skill/scripts/check_layout.py','broken.pptx'],{encoding:'utf8'});
 assert.strictEqual(failed.status,2);
 assert(JSON.parse(failed.stdout).errors.some(e=>e.code==='collapsed_text_geometry'));
 await build('report.pptx',good);
 const passed=spawnSync('python',['/skill/scripts/check_layout.py','report.pptx','--no-images'],{encoding:'utf8'});
 assert.strictEqual(passed.status,0,passed.stdout+passed.stderr);
 require('fs').writeFileSync('qa.json',passed.stdout);
 console.log('Dropped text options detected; corrected Chinese geometry passed; private QA retained.');
})().catch(e=>{console.error(e);process.exit(1)});
"""
        self.prepare(request(code, "node", ["report.pptx"]))
        self.run_job()
        self.assertEqual(read(self.job / "state.json")["status"], "completed")
        result = read(self.job / "result.json")
        self.assertEqual([f['filename'] for f in result['files']], ['report.pptx'])
        self.assertIn('corrected Chinese geometry passed', read(self.job / 'execution.json')['log'])

    def test_real_long_documents_reuse_saved_sections_and_validate_all_pages(self):
        """Multi-step authoring, actual Office/PDF files, no host or model shortcut."""
        workspaces = Workspaces(self.root)
        image = os.environ["WORKSPACE_IMAGE_ID"]
        handle = workspaces.create(runtime_image=image)["workspace_id"]
        write(self.root / "worker.json", {"image":image})
        write(self.root / "task-code.json", {"enabled":True,"image":image})
        sections = [{"title":f"Section {i+1}: 测试报告", "body":f"EVIDENCE-{i+1:02d} 已保存的分析内容，仅用于文件生成回归测试。"} for i in range(12)]
        # Save bounded chunks, as an author can do before final rendering.
        revision = 0
        for part in range(3):
            saved = workspaces.edit(handle, revision, [{"path":f"part-{part}.json", "content":json.dumps(sections[part*4:(part+1)*4],ensure_ascii=False)}], [], [])
            revision = saved["revision"]
        source_hashes = dict(workspaces.state(handle)["files"])
        code = """import json
from pathlib import Path
from pptx import Presentation
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from pypdf import PdfReader
sections = sum([json.loads(Path(f'part-{i}.json').read_text()) for i in range(3)], [])
assert len(sections)==12
p=Presentation(); d=Document()
pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
c=canvas.Canvas('report.pdf')
for i, section in enumerate(sections):
 s=p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text=section['title']; s.placeholders[1].text=section['body']
 if i: d.add_page_break()
 d.add_heading(section['title']); d.add_paragraph(section['body'])
 c.setFont('STSong-Light',18); c.drawString(40,790,section['title']); c.setFont('STSong-Light',12); c.drawString(40,730,section['body']); c.showPage()
p.save('report.pptx'); d.save('report.docx'); c.save()
assert len(Presentation('report.pptx').slides)==12
assert sum(p.text.startswith('EVIDENCE-') for p in Document('report.docx').paragraphs)==12
pdf=PdfReader('report.pdf'); assert len(pdf.pages)==12
for i,page in enumerate(pdf.pages): assert f'EVIDENCE-{i+1:02d}' in page.extract_text()
print('PPTX 12 slides, DOCX 12 sections, PDF 12 pages: content verified')
"""
        data = {"workspace_id":handle,"workspace_revision":revision,"runtime_image":image,
                "package_name":None,"package_digest":None,"runtime":"python","entrypoint":"main.py",
                "files":[{"path":"main.py","content":code}],"outputs":["report.pptx","report.docx","report.pdf"]}
        self.prepare(data)
        state=workspaces.state(handle, revision)
        manifest=[]
        for i,(name,item) in enumerate(state["files"].items()):
            target=self.job/"input"/f"artifact-{i}"
            target.write_bytes(workspaces.bytes(handle,name,state)); target.chmod(0o644)
            manifest.append({"file":target.name,"path":name,**item})
        write(self.job/"input/files.json",manifest); (self.job/"input/files.json").chmod(0o644)
        state["active_job"]=self.job.name
        write(workspaces.directory(handle)/"state.json",state)
        self.run_job()
        self.assertEqual(read(self.job/"state.json")["status"],"completed")
        self.assertIn("content verified",read(self.job/"execution.json")["log"])
        for name,sha in source_hashes.items(): self.assertEqual(workspaces.state(handle)["files"][name],sha)
        for name in ["report.pptx","report.docx"]:
            with zipfile.ZipFile(self.job/"published"/name) as z: self.assertIsNone(z.testzip())
        self.assertTrue((self.job/"published/report.pdf").read_bytes().startswith(b"%PDF"))

    def test_real_docx_skill_node_creation_pandoc_readback_and_writer_pdf(self):
        # Exercise the DOCX Skill's actual default library, not only python-docx.
        original = Path(__file__).resolve().parents[3] / "skills/pptx/scripts/office"
        shutil.copytree(original, self.package/"office", dirs_exist_ok=True)
        code = """const fs=require('fs');
const {Document,Packer,Paragraph,TextRun}=require('docx');
const {execFileSync}=require('child_process');
(async()=>{
const children=[];
for(let i=1;i<=3;i++){
 children.push(new Paragraph({text:'文档生成验证 '+i,pageBreakBefore:i>1}));
 children.push(new Paragraph({children:[new TextRun('EVIDENCE-0'+i+' 中文段落、机会与风险，仅用于回归测试。')]}));
}
const doc=new Document({styles:{default:{document:{run:{font:'Noto Sans CJK SC',size:24}}}},sections:[{children}]});
fs.writeFileSync('report.docx',await Packer.toBuffer(doc));
execFileSync('pandoc',['report.docx','-t','plain','-o','readback.txt'],{timeout:60000});
const text=fs.readFileSync('readback.txt','utf8');
for(let i=1;i<=3;i++)if(!text.includes('EVIDENCE-0'+i))throw new Error('Word text missing');
console.log(execFileSync('python',['/skill/office/soffice.py','--headless','--convert-to','pdf','--outdir','out','report.docx'],{timeout:60000}).toString());
console.log(execFileSync('python',['-c',"from pypdf import PdfReader; p=PdfReader('out/report.pdf'); assert len(p.pages)==3; assert all('EVIDENCE-0'+str(i+1) in page.extract_text() for i,page in enumerate(p.pages)); print('DOCX Node + Pandoc + Writer PDF: 3 pages verified')"],{timeout:60000}).toString());
})();"""
        self.prepare(request(code,"node",["report.docx","out/report.pdf","readback.txt"]))
        self.run_job()
        self.assertIn("3 pages verified",read(self.job/"execution.json")["log"])
        self.assertEqual(len(read(self.job/"result.json")["files"]),3)

    def test_real_followup_uses_staged_input_without_regenerating(self):
        self.prepare(request("from pathlib import Path\nPath('output.txt').write_text('immutable original')"))
        self.run_job()
        original = (self.job / "published/output.txt").read_bytes()
        self.job = self.root / "jobs" / ("work_" + os.urandom(32).hex())
        self.prepare(request("from pathlib import Path\nassert Path('previous/data.txt').read_text() == 'immutable original'\nPath('output.txt').write_text('checked original')\nprint('QA: reused prior file')"))
        (self.job / "input/artifact-0").write_bytes(original)
        write(self.job / "input/files.json", [{"file": "artifact-0", "path": "previous/data.txt", "size_bytes": len(original), "sha256": hashlib.sha256(original).hexdigest()}])
        (self.job / "input/files.json").chmod(0o644)
        self.run_job()
        self.assertEqual((self.job / "published/output.txt").read_text(), "checked original")
        self.assertIn("QA: reused", read(self.job / "execution.json")["log"])

    def test_real_inspection_only_returns_log_without_fake_output_file(self):
        data = request("print('QA: all checks passed')")
        data["outputs"] = []
        self.prepare(data)
        self.run_job()
        self.assertEqual(read(self.job / "result.json")["files"], [])
        self.assertIn("all checks passed", read(self.job / "execution.json")["log"])

    def test_real_image_qa_files_persist_without_becoming_deliverables(self):
        """Saving images is not publishing them; explicit later export stays valid."""
        workspaces = Workspaces(self.root)
        image = os.environ["WORKSPACE_IMAGE_ID"]
        handle = workspaces.create(runtime_image=image)["workspace_id"]
        write(self.root / "worker.json", {"image": image})
        write(self.root / "task-code.json", {"enabled": True, "image": image})
        steps = [
            ("from PIL import Image\nim=Image.new('RGB',(120,180),'navy')\n"
             "im.save('poster.png')\nim.crop((0,0,120,40)).save('qa_title.png')", []),
            ("from PIL import Image\nim=Image.open('poster.png')\n"
             "assert im.size==(120,180)\nassert Image.open('qa_title.png').size==(120,40)\n"
             "im.crop((0,40,120,100)).save('qa_mid.png')\nprint('private QA complete')", []),
            ("from PIL import Image\nImage.open('poster.png').verify()", ["poster.png"]),
            # A user can explicitly ask for a crop; no magic filename filtering.
            ("from PIL import Image\nImage.open('qa_title.png').verify()", ["qa_title.png"]),
        ]
        poster_sha = None
        for revision, (code, outputs) in enumerate(steps):
            if revision:
                self.job = self.root / "jobs" / ("work_" + os.urandom(32).hex())
            data = {"workspace_id": handle, "workspace_revision": revision,
                    "runtime_image": image, "package_name": None, "package_digest": None,
                    "runtime": "python", "entrypoint": f"step-{revision}.py",
                    "files": [{"path": f"step-{revision}.py", "content": code}], "outputs": outputs}
            self.prepare(data)
            state = workspaces.state(handle, revision)
            manifest = []
            for i, (name, item) in enumerate(state["files"].items()):
                target = self.job / "input" / f"artifact-{i}"
                target.write_bytes(workspaces.bytes(handle, name, state)); target.chmod(0o644)
                manifest.append({"file": target.name, "path": name, **item})
            write(self.job / "input/files.json", manifest)
            (self.job / "input/files.json").chmod(0o644)
            state["active_job"] = self.job.name
            write(workspaces.directory(handle) / "state.json", state)
            self.run_job()
            self.assertEqual(read(self.job / "state.json")["status"], "completed")
            self.assertEqual([f["filename"] for f in read(self.job / "result.json")["files"]], outputs)
            current_sha = hashlib.sha256(workspaces.bytes(handle, "poster.png")).hexdigest()
            if poster_sha is None: poster_sha = current_sha
            self.assertEqual(current_sha, poster_sha, "QA and delivery must reuse the original image")
            self.assertTrue(workspaces.bytes(handle, "qa_title.png").startswith(b'\x89PNG'))
            if revision: self.assertTrue(workspaces.bytes(handle, "qa_mid.png").startswith(b'\x89PNG'))

    def test_real_factorial_stdout_has_no_required_artifact(self):
        import math
        data = request("import math\nprint(math.factorial(100))")
        data["outputs"] = []
        self.prepare(data)
        self.run_job()
        self.assertEqual(read(self.job / "state.json")["status"], "completed")
        self.assertEqual(read(self.job / "result.json")["files"], [])
        execution = read(self.job / "execution.json")
        self.assertEqual(execution["exit_code"], 0)
        self.assertEqual(execution["log"].strip(), str(math.factorial(100)))

    def test_real_missing_declared_file_keeps_computed_stdout(self):
        import math
        self.prepare(request("import math\nprint(math.factorial(100))", outputs=["answer.txt"]))
        with self.assertRaisesRegex(ValueError, "output missing.*answer.txt"):
            self.run_job()
        execution = read(self.job / "execution.json")
        self.assertIn(str(math.factorial(100)), execution["log"])
        self.assertEqual(execution["exit_code"], 1)
        self.assertFalse((self.job / "published/answer.txt").exists())

    def test_real_cancellation_terminates_container(self):
        self.prepare(request("import time; time.sleep(200)"))
        errors = []
        def execute():
            try: self.run_job()
            except Exception as error: errors.append(error)
        thread = threading.Thread(target=execute)
        thread.start()
        deadline = time.time() + 30
        while time.time() < deadline and read(self.job / "state.json")["status"] != "working":
            time.sleep(0.2)
        (self.job / "cancel").touch()
        thread.join(30)
        self.assertFalse(thread.is_alive())
        self.assertEqual(errors, [])
        self.assertEqual(read(self.job / "state.json")["status"], "cancelled")

    def test_real_workspace_python_plot_xlsx_failure_patch_and_script_arguments(self):
        workspaces = Workspaces(self.root)
        handle = workspaces.create()["workspace_id"]
        write(self.root / "worker.json", {"image":os.environ["WORKSPACE_IMAGE_ID"]})
        write(self.root / "task-code.json", {"enabled":True, "image":os.environ["WORKSPACE_IMAGE_ID"]})
        code = """from pathlib import Path
import sys, os, socket
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook
assert os.getuid() == 65532
assert not Path('/var/run/docker.sock').exists()
assert not Path('/run/secrets/workspace_mcp_token').exists()
try:
 socket.create_connection(('1.1.1.1',443),timeout=1)
 raise AssertionError('network was not isolated')
except OSError: pass
data = pd.DataFrame({'year':[2020,2021,2022,2023,2024], 'value':[1,3,2,5,4]})
data.to_csv('data.csv',index=False)
assert False, 'deliberate local bug'
assert sys.argv[1:] == ['a b','; touch /tmp/not-a-shell-command']
assert not Path('/tmp/not-a-shell-command').exists()
saved = pd.read_csv('data.csv')
plt.plot(saved['year'],saved['value']); plt.savefig('chart.png')
book=Workbook(); sheet=book.active
for row in [list(saved.columns), *saved.values.tolist()]: sheet.append(row)
book.save('report.xlsx')
"""
        draft = workspaces.edit(handle, 0, [{"path":"main.py", "content":code}], [], [])
        first_hash = draft["files"][0]["sha256"]

        def prepare_workspace(revision):
            data = {"workspace_id":handle, "workspace_revision":revision, "package_name":None, "package_digest":None,
                    "runtime":"python", "entrypoint":"main.py", "entrypoint_root":"workspace", "files":[],
                    "arguments":["a b", "; touch /tmp/not-a-shell-command"], "outputs":["chart.png","report.xlsx"]}
            self.prepare(data)
            state = workspaces.state(handle, revision)
            manifest = []
            for i, (name, item) in enumerate(state["files"].items()):
                path = self.job / "input" / f"artifact-{i}"
                path.write_bytes(workspaces.bytes(handle, name, state)); path.chmod(0o644)
                manifest.append({"file":path.name,"path":name, **item})
            write(self.job / "input/files.json", manifest)
            (self.job / "input/files.json").chmod(0o644)
            state["active_job"] = self.job.name
            write(workspaces.directory(handle) / "state.json", state)

        prepare_workspace(1)
        with self.assertRaisesRegex(ValueError, "deliberate local bug"):
            self.run_job()
        write(self.job / "state.json", {"status":"failed"})
        failed = workspaces.inspect(handle)
        self.assertEqual(failed["revision"], 2)
        csv = workspaces.bytes(handle,"data.csv")
        patched = workspaces.edit(handle, 2, [], [{"path":"main.py", "sha256":first_hash,
            "before":"assert False, 'deliberate local bug'", "after":"pass  # repaired only this assertion"}], [])
        self.assertEqual(workspaces.bytes(handle,"data.csv"), csv)
        first = self.job.name
        self.job = self.root / "jobs" / ("work_" + os.urandom(32).hex())
        prepare_workspace(patched["revision"])
        self.run_job()
        self.assertEqual(read(self.job / "state.json")["status"], "completed")
        self.assertTrue((self.job / "published/chart.png").read_bytes().startswith(b"\x89PNG"))
        with zipfile.ZipFile(self.job / "published/report.xlsx") as archive:
            self.assertIn("xl/workbook.xml", archive.namelist())
        self.assertEqual(workspaces.inspect(handle)["revision"], 4)
        self.assertEqual(workspaces.bytes(handle,"data.csv"), csv)
        docker("rm", "-f", "exoharness-" + first, check=False)

    def test_real_registered_skill_script_receives_argument_array(self):
        code = "from pathlib import Path\nimport sys\nPath('output.txt').write_text('|'.join(sys.argv[1:]))"
        (self.package / "check.py").write_text(code)
        self.package.chmod(0o700)
        (self.package / "check.py").chmod(0o600)
        data = request()
        data.update(entrypoint_root="skill", entrypoint="check.py", files=[{"path":"check.py","content":"raise Exception('must use the registered script')"}], arguments=["hello world","--flag"])
        self.prepare(data)
        self.run_job()
        self.assertEqual((self.job / "published/output.txt").read_text(), "hello world|--flag")
        self.assertEqual(self.package.stat().st_mode & 0o777, 0o700)
        self.assertEqual((self.package / "check.py").stat().st_mode & 0o777, 0o600)
        self.assertEqual((self.job / "skill/check.py").read_text(), code)

    def test_real_pptx_shape_failure_repairs_only_saved_code_and_reuses_content(self):
        image = os.environ["WORKSPACE_IMAGE_ID"]
        write(self.root / "worker.json", {"image": image})
        write(self.root / "inventory.json", {"image": image, "commands": {"node": True}})
        write(self.root / "task-code.json", {"enabled": True, "image": image})
        workspaces = Workspaces(self.root)
        workspace = workspaces.create(runtime_image=image)
        handle = workspace["workspace_id"]
        content = json.dumps(["SECTION-1", "SECTION-2", "SECTION-END"])
        code = """const fs = require('fs');
const PptxGenJS = require('pptxgenjs');
const p = new PptxGenJS(); p.layout = 'LAYOUT_WIDE';
for (const text of JSON.parse(fs.readFileSync('content.json','utf8'))) {
 const s=p.addSlide();
 s.addShape(p.ShapeType.oval,{x:1,y:1,w:1,h:1});
 s.addText(text,{x:2,y:2,w:8,h:1});
}
p.writeFile({fileName:'report.pptx'});
"""
        state = workspaces.edit(handle, 0, [{"path": "content.json", "content": content}, {"path": "build.js", "content": code}], [], [])
        original = {f["path"]: f["sha256"] for f in state["files"]}
        def submit(revision):
            self.job = self.root / "jobs" / ("work_" + os.urandom(32).hex())
            self.prepare({"workspace_id": handle, "workspace_revision": revision,
                "package_name": None, "package_digest": None, "runtime_image": image,
                "runtime": "node", "entrypoint": "build.js", "entrypoint_root": "workspace",
                "files": [], "arguments": [], "outputs": ["report.pptx"]})
            current = workspaces.state(handle, revision)
            inputs = []
            for index, (name, metadata) in enumerate(current["files"].items()):
                path = self.job / "input" / f"artifact-{index}"
                path.write_bytes(workspaces.bytes(handle, name, current)); path.chmod(0o644)
                inputs.append({"file": path.name, "path": name, **metadata})
            write(self.job / "input/files.json", inputs)
            (self.job / "input/files.json").chmod(0o644)
            current["active_job"] = self.job.name
            write(workspaces.directory(handle) / "state.json", current)
        submit(state["revision"])
        with self.assertRaisesRegex(ValueError, "Missing/Invalid shape"):
            self.run_job()
        write(self.job / "state.json", {"status": "failed"})
        failed_job = self.job
        state = workspaces.inspect(handle)
        patched = workspaces.edit(handle, state["revision"], [], [{"path": "build.js", "sha256": original["build.js"],
            "before": "p.ShapeType.oval", "after": "p.ShapeType.ellipse"}], [])
        self.assertEqual(workspaces.bytes(handle, "content.json"), content.encode())
        self.assertEqual(workspaces.bytes(handle, "build.js"), code.replace("p.ShapeType.oval", "p.ShapeType.ellipse").encode())
        submit(patched["revision"])
        self.assertEqual(read(self.job / "request.json")["files"], [])  # no full code/content resubmission
        self.run_job()
        with zipfile.ZipFile(self.job / "published/report.pptx") as archive:
            slides = [n for n in archive.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
            self.assertEqual(len(slides), 3)
            for idx, marker in enumerate(json.loads(content), 1):
                self.assertIn(marker.encode(), archive.read(f"ppt/slides/slide{idx}.xml"))
        self.assertEqual(workspaces.bytes(handle, "content.json"), content.encode())
        self.assertEqual(read(failed_job / "request.json")["files"], [])

    def test_inventory_is_measured_inside_the_fixed_image(self):
        inventory = runtime_inventory(os.environ["WORKSPACE_IMAGE_ID"])
        self.assertEqual(inventory["python_packages"]["matplotlib"], "3.10.6")
        self.assertEqual(inventory["python_packages"]["openpyxl"], "3.1.5")
        self.assertTrue(inventory["commands"]["python"])
        self.assertTrue(inventory["commands"]["bash"])

    def test_cjk_plotting_defaults_render_without_missing_glyphs(self):
        self.prepare(request("import matplotlib.pyplot as plt\nplt.plot([1,2],[3,4]); plt.title('中文图表：年份与数值'); plt.xlabel('年份'); plt.savefig('chart.png')", outputs=["chart.png"]))
        self.run_job()
        self.assertTrue((self.job / "published/chart.png").read_bytes().startswith(b"\x89PNG"))
        self.assertNotIn("missing from font", read(self.job / "execution.json")["log"])

    def test_missing_glyph_plot_is_an_error_not_a_published_artifact(self):
        self.prepare(request("import matplotlib.pyplot as plt\nplt.rcParams['font.sans-serif']=['DejaVu Sans']\nplt.title('中文图表'); plt.savefig('broken.png')", outputs=["broken.png"]))
        with self.assertRaisesRegex(ValueError, "Glyph"):
            self.run_job()
        self.assertFalse((self.job / "result.json").exists())
        self.assertFalse((self.job / "published/broken.png").exists())

    def test_shell_pipeline_and_python_plot_share_task_sandbox_without_skill_binding(self):
        workspaces = Workspaces(self.root)
        workspace = workspaces.create(runtime_image=os.environ["WORKSPACE_IMAGE_ID"])
        write(self.root / "worker.json", {"image": os.environ["WORKSPACE_IMAGE_ID"]})
        write(self.root / "task-code.json", {"enabled": True, "image": os.environ["WORKSPACE_IMAGE_ID"]})
        write(self.root / "packages.json", {})
        code = """set -euo pipefail
test "$(id -u)" = 65532
test ! -e /var/run/docker.sock
test ! -d /run/secrets
test ! -e /mnt/data/ExoHarness
if touch /etc/sandbox-write-probe 2>/dev/null; then exit 99; fi
test "$1" = '; touch /tmp/should-not-exist'
test ! -e /tmp/should-not-exist
printf 'group,value\\nA,1\\nB,4\\nC,2\\n' | tee rows.csv >/dev/null
python - <<'PY'
from pathlib import Path
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
data = pd.read_csv('rows.csv')
plt.plot(data['group'], data['value']); plt.savefig('chart.png')
assert not Path('/mnt/data/ExoHarness').exists()
import socket
try:
 socket.create_connection(('1.1.1.1', 443), timeout=1)
 raise AssertionError('Shell subprocess escaped network isolation')
except OSError: pass
PY
"""
        data = {"workspace_id": workspace["workspace_id"], "workspace_revision": 0,
                "runtime_image": os.environ["WORKSPACE_IMAGE_ID"], "package_name": None, "package_digest": None,
                "runtime": "shell", "entrypoint": "script.sh", "entrypoint_root": "workspace",
                "files": [{"path": "script.sh", "content": code}], "arguments": ["; touch /tmp/should-not-exist"],
                "outputs": ["rows.csv", "chart.png"]}
        self.prepare(data)
        state = workspaces.state(workspace["workspace_id"])
        state["active_job"] = self.job.name
        write(workspaces.directory(workspace["workspace_id"]) / "state.json", state)
        self.run_job()
        self.assertEqual(read(self.job / "state.json")["status"], "completed")
        self.assertTrue((self.job / "published/chart.png").read_bytes().startswith(b"\x89PNG"))
        self.assertEqual(workspaces.bytes(workspace["workspace_id"], "rows.csv"), b"group,value\nA,1\nB,4\nC,2\n")

    def test_real_xlsx_formulas_are_recalculated_not_just_written(self):
        code = """from openpyxl import Workbook, load_workbook
import subprocess
from pathlib import Path
book=Workbook(); sheet=book.active
sheet['A1']=2; sheet['A2']=3; sheet['A3']='=SUM(A1:A2)'
book.save('input.xlsx')
process=subprocess.run(['libreoffice','--headless','-env:UserInstallation=file:///tmp/calc-profile','--convert-to','xlsx','--outdir','recalculated','input.xlsx'],capture_output=True,text=True,timeout=60)
assert process.returncode == 0, process.stderr
result=Path('recalculated/input.xlsx')
assert result.is_file(), process.stdout + process.stderr
assert load_workbook(result,data_only=True).active['A3'].value == 5
Path('output.txt').write_text('Formula recalculated: 5')
"""
        self.prepare(request(code))
        self.run_job()
        self.assertEqual((self.job/'published/output.txt').read_text(), 'Formula recalculated: 5')


if __name__ == "__main__":
    unittest.main()
