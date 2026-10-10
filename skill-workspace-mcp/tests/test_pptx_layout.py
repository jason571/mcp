"""Regression for valid-but-unreadable PPTX and image-free deliverables."""
import hashlib
import importlib.util
import json
import subprocess
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

source = Path(__file__).resolve().parents[3] / "skills/pptx/scripts/check_layout.py"
spec = importlib.util.spec_from_file_location("pptx_layout", source)
layout = importlib.util.module_from_spec(spec)
spec.loader.exec_module(layout)


class LayoutTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / "deck.pptx"
        self.deck = Presentation()
        self.slide = self.deck.slides.add_slide(self.deck.slide_layouts[6])

    def text(self, x, y, w=2, h=.4, shapes=None):
        shape = (shapes if shapes is not None else self.slide.shapes).add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.text = "中文内容：结论与具体依据"
        return shape

    def check(self, no_images=False):
        self.deck.save(self.path)
        before = hashlib.sha256(self.path.read_bytes()).hexdigest()
        result = layout.inspect(self.path, no_images)
        self.assertEqual(before, hashlib.sha256(self.path.read_bytes()).hexdigest())
        return result

    def test_dropped_options_collapse_is_rejected_even_if_pptx_opens(self):
        for _ in range(8): self.text(0, 0, 3, 3)
        result = self.check()
        self.assertEqual(result['status'], 'failed')
        self.assertEqual(result['errors'][0]['code'], 'collapsed_text_geometry')

    def test_spaced_chinese_text_passes_without_editing(self):
        for i in range(5): self.text(.6, .5 + i * .8, 7, .6)
        self.assertEqual(self.check()['status'], 'passed')

    def test_replacement_character_in_chinese_text_is_rejected(self):
        shape = self.text(.6, .5, 7, .6)
        shape.text = '行业现状：营收增\ufffd'
        result = self.check()
        self.assertEqual(result['status'], 'failed')
        self.assertEqual(result['errors'][0]['code'], 'replacement_character_in_text')

    def test_broken_encoding_in_table_cell_is_rejected(self):
        table = self.slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(4), Inches(1)).table
        table.cell(0, 0).text = '增长\x81趋势'
        result = self.check()
        self.assertEqual(result['errors'][0]['code'], 'control_character_in_text')
        self.assertEqual(result['errors'][0]['table_cell'], [1, 1])

    def test_reversible_chinese_mojibake_is_rejected(self):
        shape = self.text(.6, .5, 7, .6)
        shape.text = 'ä¸­æ–‡æ ‡é¢˜'
        self.assertEqual(self.check()['errors'][0]['code'], 'mojibake_chinese_text')

    def test_valid_chinese_table_text_passes(self):
        table = self.slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(4), Inches(1)).table
        table.cell(0, 0).text = '行业现状：增长趋势'
        self.assertEqual(self.check()['status'], 'passed')

    def test_two_intentionally_layered_labels_are_not_mass_collapse(self):
        self.text(1, 1); self.text(1, 1)
        self.assertEqual(self.check()['status'], 'passed')

    def test_offslide_and_zero_dimension_text_are_reported(self):
        self.text(9.8, 1); self.text(1, 2, 0, .4)
        codes = {e['code'] for e in self.check()['errors']}
        self.assertEqual(codes, {'text_outside_slide', 'invalid_text_dimensions'})

    def test_group_coordinates_are_respected(self):
        group = self.slide.shapes.add_group_shape()
        self.text(1, 1, shapes=group.shapes)
        self.text(1, 2, shapes=group.shapes)
        group.left, group.top = Inches(3), Inches(2)
        self.assertEqual(self.check()['status'], 'passed')

    def test_images_are_only_rejected_on_explicit_no_image_check(self):
        pic = Path(self.tmp.name) / 'sample.png'
        Image.new('RGB', (10, 10), 'red').save(pic)
        self.slide.shapes.add_picture(str(pic), Inches(1), Inches(1))
        self.assertEqual(self.check()['status'], 'passed')
        self.assertEqual(self.check(True)['errors'][0]['code'], 'image_not_requested')

    def finalize(self, no_images=False):
        self.deck.save(self.path)
        target = Path(self.tmp.name) / 'delivered.pptx'
        script = """const fs = require('node:fs');
const finalize = require(process.argv[1]);
finalize({writeFile: async ({fileName}) => fs.copyFileSync(process.argv[2], fileName)},
  process.argv[3], {noImages: process.argv[4] === 'true'})
  .catch(error => { console.error(error.message); process.exitCode = 1; });"""
        run = subprocess.run(['node', '-e', script, str(source.with_name('finalize.cjs')),
                              str(self.path), str(target), str(no_images).lower()],
                             capture_output=True, text=True, timeout=150)
        return run, target

    def test_finalizer_publishes_only_after_actual_checks(self):
        self.text(.6, .5, 7, .6)
        run, target = self.finalize()
        self.assertEqual(run.returncode, 0, run.stderr)
        self.assertTrue(target.is_file())
        evidence = json.loads(run.stdout)
        self.assertEqual(evidence['slides'], 1)
        self.assertEqual(evidence['checks'], ['office_structure', 'text_geometry', 'text_integrity', 'table_text_overlap', 'cjk_table_portability'])
        self.assertFalse(target.with_name('.delivered.pptx.pending.pptx').exists())

    def test_finalizer_rejects_real_offslide_result_without_publishing(self):
        # Exact first geometry error from the real four-slide regression deck.
        self.text(.79, 6.78, 5.517, .78)
        run, target = self.finalize()
        self.assertNotEqual(run.returncode, 0)
        self.assertIn('text_outside_slide', run.stderr)
        self.assertFalse(target.exists())
        self.assertTrue(target.with_name('.delivered.pptx.pending.pptx').exists())
        self.assertTrue(self.path.exists())

    def test_finalizer_does_not_publish_garbled_chinese(self):
        shape = self.text(.6, .5, 7, .6)
        shape.text = '中文标题\ufffd'
        run, target = self.finalize()
        self.assertNotEqual(run.returncode, 0)
        self.assertIn('replacement_character_in_text', run.stderr)
        self.assertFalse(target.exists())

    def test_finalizer_failed_replacement_retains_previous_delivery(self):
        target = Path(self.tmp.name) / 'delivered.pptx'
        previous = b'previous delivery remains unchanged'
        target.write_bytes(previous)
        self.text(1, 7.4)
        run, target = self.finalize()
        self.assertNotEqual(run.returncode, 0)
        self.assertEqual(target.read_bytes(), previous)

    def test_finalizer_honors_optional_image_policy(self):
        pic = Path(self.tmp.name) / 'sample.png'
        Image.new('RGB', (10, 10), 'red').save(pic)
        self.slide.shapes.add_picture(str(pic), Inches(1), Inches(1))
        rejected, target = self.finalize(no_images=True)
        self.assertNotEqual(rejected.returncode, 0)
        self.assertIn('image_not_requested', rejected.stderr)
        self.assertFalse(target.exists())
        accepted, target = self.finalize()
        self.assertEqual(accepted.returncode, 0, accepted.stderr)
        self.assertTrue(target.is_file())

    def test_new_deck_table_overlap_blocks_publish_without_changing_template_checks(self):
        table = self.slide.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(6), Inches(2))
        table.table.cell(0, 0).text = '数据'
        self.text(1.1, 2.8, 5, .6)
        self.deck.save(self.path)
        self.assertEqual(layout.inspect(self.path)['status'], 'passed')
        error = layout.inspect(self.path, table_overlap=True)['errors'][0]
        self.assertEqual(error['code'], 'text_overlaps_table')
        run, target = self.finalize()
        self.assertNotEqual(run.returncode, 0)
        self.assertIn('text_overlaps_table', run.stderr)
        self.assertFalse(target.exists())

    def test_table_caption_below_table_is_not_rejected(self):
        self.slide.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(6), Inches(2))
        self.text(1, 3.1, 6, .5)
        run, target = self.finalize()
        self.assertEqual(run.returncode, 0, run.stderr)
        self.assertTrue(target.is_file())

    def test_new_deck_rejects_cjk_table_with_runtime_only_font(self):
        table = self.slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(6), Inches(1)).table
        run = table.cell(0, 0).text_frame.paragraphs[0].add_run()
        run.text = '千年文脉'
        run.font.name = 'Noto Sans CJK SC'
        run._r.rPr.set('lang', 'en-US')
        self.deck.save(self.path)
        self.assertEqual(layout.inspect(self.path)['status'], 'passed')
        codes = {error['code'] for error in layout.inspect(self.path, table_overlap=True)['errors']}
        self.assertEqual(codes, {'cjk_table_font_not_portable', 'cjk_table_language_mismatch'})
        rejected, target = self.finalize()
        self.assertNotEqual(rejected.returncode, 0)
        self.assertFalse(target.exists())

    def test_new_deck_accepts_explicit_target_font_and_language(self):
        table = self.slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(6), Inches(1)).table
        run = table.cell(0, 0).text_frame.paragraphs[0].add_run()
        run.text = '千年文脉'
        run.font.name = 'Microsoft YaHei'
        run._r.rPr.set('lang', 'zh-CN')
        self.assertEqual(self.check()['status'], 'passed')
        self.deck.save(self.path)
        self.assertEqual(layout.inspect(self.path, table_overlap=True)['status'], 'passed')

    def test_real_generator_stub_height_does_not_hide_table_overlap(self):
        self.deck.slide_width = Inches(13.333)
        table = self.slide.shapes.add_table(6, 3, Inches(7.03), Inches(2.68), Inches(5.48), Inches(3.36))
        table.height = Inches(1)  # Actual failed pptxgenjs output used this stub.
        self.text(7.03, 5.68, 5.48, .86)
        self.deck.save(self.path)
        self.assertEqual(layout.inspect(self.path)['status'], 'passed')
        error = layout.inspect(self.path, table_overlap=True)['errors'][0]
        self.assertEqual(error['code'], 'text_overlaps_table')
        self.assertEqual(error['table_box_inches'], [7.03, 2.68, 5.48, 3.36])


if __name__ == '__main__': unittest.main()
